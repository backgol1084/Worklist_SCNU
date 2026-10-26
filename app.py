import streamlit as st
from supabase import create_client, Client
import pandas as pd
from datetime import datetime, date
import openpyxl
from io import BytesIO

# 1. 페이지 기본 설정 (가장 처음에 위치해야 함)
st.set_page_config(
    page_title="HYO'S 업무 진행 상황 관리",
    page_icon="✅",
    layout="wide"
)

# 2. Supabase 클라이언트 초기화
@st.cache_resource
def init_client():
    url = st.secrets["SUPABASE_URL"]
    key = st.secrets["SUPABASE_KEY"]
    return create_client(url, key)

supabase: Client = init_client()

# 3. 데이터베이스 연동 함수들
def load_data():
    """DB에서 전체 업무 리스트 불러오기"""
    try:
        response = supabase.table("tasks").select("*").order("id", desc=False).execute()
        return response.data
    except Exception as e:
        st.error(f"데이터를 불러오는 중 오류가 발생했습니다: {e}")
        return []

def add_task_db(category, title, subtasks, start_date, due_date):
    """새로운 업무, 세부 체크리스트 및 일정 추가 (세부 업무별 날짜 기본값 포함)"""
    new_row = {
        "category": category,
        "title": title,
        "status": "대기",
        "subtasks": [{"name": item, "done": False, "memo": "", "start_date": str(start_date) if start_date else None, "end_date": str(due_date) if due_date else None} for item in subtasks],
        "start_date": str(start_date) if start_date else None,
        "due_date": str(due_date) if due_date else None,
        "end_date": None
    }
    supabase.table("tasks").insert(new_row).execute()

def update_task_db(task_id, subtasks, status, start_date=None, due_date=None, end_date=None):
    """업무 정보, 세부 업무, 상태 및 일정 업데이트"""
    update_data = {
        "subtasks": subtasks,
        "status": status
    }
    if start_date is not None:
        update_data["start_date"] = str(start_date) if start_date else None
    if due_date is not None:
        update_data["due_date"] = str(due_date) if due_date else None
    if end_date is not None:
        update_data["end_date"] = str(end_date) if end_date else None

    supabase.table("tasks").update(update_data).eq("id", task_id).execute()

def delete_task_db(task_id):
    """업무 삭제"""
    supabase.table("tasks").delete().eq("id", task_id).execute()

def _update_and_save(task_id, subtasks):
    """진행률 및 상태 자동 계산 후 DB 저장 헬퍼 함수 (완료 시 종료일 자동 기록)"""
    total = len(subtasks)
    done = sum(1 for s in subtasks if s["done"])
    prog = (done / total) if total > 0 else 0
    
    tasks = load_data()
    curr = next((t for t in tasks if t["id"] == task_id), None)
    current_end_date = curr.get("end_date") if curr else None

    if prog == 1.0:
        status = "완료"
        if not current_end_date:
            current_end_date = str(date.today())
    elif prog > 0:
        status = "진행중"
        current_end_date = None
    else:
        status = "대기"
        current_end_date = None

    update_task_db(task_id, subtasks, status, end_date=current_end_date)

def _clear_widget_keys(task_id):
    """위젯 충돌 방지를 위해 해당 업무의 세션 상태 키 모두 초기화"""
    keys_to_del = [k for k in st.session_state.keys() if k.startswith(f"edit_chk_{task_id}_") or k.startswith(f"edit_memo_{task_id}_") or k.startswith(f"edit_sub_start_{task_id}_") or k.startswith(f"edit_sub_end_{task_id}_")]
    for k in keys_to_del:
        del st.session_state[k]

# 세션 상태에 편집 중인 업무 ID 관리 초기화
if "editing_task_id" not in st.session_state:
    st.session_state.editing_task_id = None

# --- 커스텀 CSS (메모 박스 크기 조절 허용 및 기본 높이 확장) ---
st.markdown("""
<style>
    textarea {
        resize: vertical !important;
        min-height: 100px !important;
    }
</style>
""", unsafe_allow_html=True)

# --- 앱 UI 레이아웃 시작 ---
st.title("✅ HYO'S 진행 상황 관리 앱")
st.caption("나만의 간편한 업무 리스트 & 실시간 클라우드 DB 트래커")
st.markdown("---")

# 4. 사이드바: 새 업무 등록 폼
with st.sidebar:
    st.header("➕ 새 업무 등록")
    with st.form("add_task_form", clear_on_submit=True):
        category = st.text_input("업무 구분", placeholder="예: 기획, 개발, 보고")
        task_title = st.text_input("업무 제목", placeholder="예: 3분기 실적 보고서 작성")
        
        col_d1, col_d2 = st.columns(2)
        with col_d1:
            start_date = st.date_input("시작일", value=date.today())
        with col_d2:
            due_date = st.date_input("마감예정일", value=date.today())

        subtasks_text = st.text_area(
            "세부 업무 리스트 (줄바꿈으로 구분)",
            placeholder="자료 수집\n초안 작성\n팀장님 검토\n최종 제출"
        )
        
        submitted = st.form_submit_button("업무 추가하기", use_container_width=True)
        
        if submitted:
            if category.strip() and task_title.strip():
                with st.spinner("새 업무를 등록하는 중입니다..."):
                    sub_items = [s.strip() for s in subtasks_text.split("\n") if s.strip()]
                    add_task_db(category.strip(), task_title.strip(), sub_items, start_date, due_date)
                st.success("새 업무가 추가되었습니다!")
                st.rerun()
            else:
                st.warning("업무 구분과 업무 제목을 모두 입력해주세요.")

# 데이터 불러오기
with st.spinner("데이터를 불러오는 중입니다..."):
    tasks = load_data()

# 5. 메인 화면: 탭 구성
tab1, tab2 = st.tabs(["📋 업무 진행 & 관리", "📊 한눈에 보기 & 엑셀 다운로드"])

with tab1:
    # --- 상단 편집 패널 ---
    if st.session_state.editing_task_id is not None:
        current_task = next((t for t in tasks if t["id"] == st.session_state.editing_task_id), None)
        if current_task:
            st.markdown(
                f"""
                <div style="padding: 20px; border-radius: 10px; border: 2px solid #4CAF50; background-color: #f1f8e9; margin-bottom: 20px;">
                    <h3 style="margin-top: 0; color: #2e7d32;">⚙️ 세부 업무, 일정 및 메모 편집 모드</h3>
                    <p style="font-size: 16px; font-weight: bold; margin-bottom: 5px;">[{current_task['category']}] {current_task['title']}</p>
                    <p style="color: #555; font-size: 13px;">전체 일정 수정, 세부 업무별 시작일/종료일 지정, 체크, 추진계획 메모 작성, 순서 변경(⬆️, ⬇️), 삭제(❌)를 자유롭게 진행하세요.</p>
                </div>
                """, 
                unsafe_allow_html=True
            )
            
            task_id = current_task["id"]
            
            # 전체 일정 수정 영역
            col_ed1, col_ed2, col_ed3 = st.columns(3)
            with col_ed1:
                cur_start = datetime.strptime(current_task["start_date"], "%Y-%m-%d").date() if current_task.get("start_date") else date.today()
                new_start = st.date_input("전체 시작일 수정", value=cur_start, key=f"edit_start_{task_id}")
            with col_ed2:
                cur_due = datetime.strptime(current_task["due_date"], "%Y-%m-%d").date() if current_task.get("due_date") else date.today()
                new_due = st.date_input("전체 마감예정일 수정", value=cur_due, key=f"edit_due_{task_id}")
            with col_ed3:
                cur_end = datetime.strptime(current_task["end_date"], "%Y-%m-%d").date() if current_task.get("end_date") else None
                new_end = st.date_input("전체 종료일 (실제 완료일)", value=cur_end, key=f"edit_end_{task_id}", help="완료된 경우 자동 입력되거나 직접 수정할 수 있습니다.")

            if new_start != cur_start or new_due != cur_due or new_end != cur_end:
                update_task_db(task_id, current_task["subtasks"], current_task["status"], start_date=new_start, due_date=new_due, end_date=new_end)

            st.markdown("---")
            st.subheader("📝 세부 업무별 체크리스트, 일정 및 추진계획 메모 관리")

            # 데이터 무결성 보정 (기존 데이터에 세부 날짜가 없는 경우 처리)
            subtasks = []
            default_sub_start = current_task.get("start_date") or str(date.today())
            default_sub_end = current_task.get("due_date") or str(date.today())

            for s in current_task["subtasks"]:
                if isinstance(s, str):
                    subtasks.append({"name": s, "done": False, "memo": "", "start_date": default_sub_start, "end_date": default_sub_end})
                else:
                    item_dict = dict(s)
                    if "memo" not in item_dict: item_dict["memo"] = ""
                    if "start_date" not in item_dict or not item_dict["start_date"]: item_dict["start_date"] = default_sub_start
                    if "end_date" not in item_dict or not item_dict["end_date"]: item_dict["end_date"] = default_sub_end
                    subtasks.append(item_dict)
            
            if not subtasks:
                st.info("등록된 세부 업무가 없습니다. 아래에서 추가해보세요.")
            
            # 세부 업무 리스트 순회 및 조작
            for idx, sub in enumerate(subtasks):
                chk_key = f"edit_chk_{task_id}_{idx}"
                memo_key = f"edit_memo_{task_id}_{idx}"
                sub_start_key = f"edit_sub_start_{task_id}_{idx}"
                sub_end_key = f"edit_sub_end_{task_id}_{idx}"
                
                if chk_key not in st.session_state:
                    st.session_state[chk_key] = sub["done"]
                if memo_key not in st.session_state:
                    st.session_state[memo_key] = sub["memo"]

                st.markdown(f"**[{idx+1}] 세부 항목 컨트롤**")
                
                # 체크박스와 순서 변경/삭제 버튼 배치
                rc1, rc2, rc3, rc4 = st.columns([5, 1, 1, 1])
                with rc1:
                    new_done = st.checkbox(sub["name"], key=chk_key)
                    if new_done != sub["done"]:
                        subtasks[idx]["done"] = new_done
                        _update_and_save(task_id, subtasks)
                with rc2:
                    if idx > 0:
                        if st.button("⬆️", key=f"edit_up_{task_id}_{idx}"):
                            subtasks[idx], subtasks[idx - 1] = subtasks[idx - 1], subtasks[idx]
                            _update_and_save(task_id, subtasks)
                            _clear_widget_keys(task_id)
                            st.rerun()
                with rc3:
                    if idx < len(subtasks) - 1:
                        if st.button("⬇️", key=f"edit_down_{task_id}_{idx}"):
                            subtasks[idx], subtasks[idx + 1] = subtasks[idx + 1], subtasks[idx]
                            _update_and_save(task_id, subtasks)
                            _clear_widget_keys(task_id)
                            st.rerun()
                with rc4:
                    if st.button("❌", key=f"edit_del_{task_id}_{idx}"):
                        subtasks.pop(idx)
                        _update_and_save(task_id, subtasks)
                        _clear_widget_keys(task_id)
                        st.rerun()

                # 세부 업무별 시작일 / 종료일 지정 컬럼
                sc1, sc2 = st.columns(2)
                with sc1:
                    cur_s_date = datetime.strptime(sub["start_date"], "%Y-%m-%d").date() if sub.get("start_date") else date.today()
                    new_sub_start = st.date_input("세부 시작일", value=cur_s_date, key=sub_start_key)
                with sc2:
                    cur_e_date = datetime.strptime(sub["end_date"], "%Y-%m-%d").date() if sub.get("end_date") else date.today()
                    new_sub_end = st.date_input("세부 종료일", value=cur_e_date, key=sub_end_key)

                if str(new_sub_start) != sub.get("start_date") or str(new_sub_end) != sub.get("end_date"):
                    subtasks[idx]["start_date"] = str(new_sub_start)
                    subtasks[idx]["end_date"] = str(new_sub_end)
                    update_task_db(task_id, subtasks, current_task["status"])

                # 추진계획 메모 입력 필드
                new_memo = st.text_area("추진계획 메모", value=sub["memo"], key=memo_key, placeholder="이 세부 업무와 관련된 추진계획이나 메모를 여러 줄로 자유롭게 입력하세요")
                if new_memo != sub["memo"]:
                    subtasks[idx]["memo"] = new_memo
                    update_task_db(task_id, subtasks, current_task["status"])

                st.markdown("---")
            
            # 새 세부 업무 추가 폼
            with st.form(key=f"edit_add_form_{task_id}", clear_on_submit=True):
                new_sub_name = st.text_input("새 세부 업무 추가", placeholder="세부 업무명 입력")
                
                acol1, acol2 = st.columns(2)
                with acol1:
                    new_sub_sdate = st.date_input("새 세부 시작일", value=date.today())
                with acol2:
                    new_sub_edate = st.date_input("새 세부 종료일", value=date.today())

                submitted_sub = st.form_submit_button("➕ 세부 업무 추가하기", use_container_width=True)
                if submitted_sub:
                    if new_sub_name.strip():
                        subtasks.append({
                            "name": new_sub_name.strip(),
                            "done": False,
                            "memo": "",
                            "start_date": str(new_sub_sdate),
                            "end_date": str(new_sub_edate)
                        })
                        _update_and_save(task_id, subtasks)
                        _clear_widget_keys(task_id)
                        st.success("세부 업무가 추가되었습니다!")
                        st.rerun()
                    else:
                        st.warning("내용을 입력해주세요.")

            st.write("")
            if st.button("📌 편집 완료 (닫기)", use_container_width=True, type="primary"):
                _clear_widget_keys(task_id)
                st.session_state.editing_task_id = None
                st.rerun()
            
            st.markdown("---")
            st.markdown("<br>", unsafe_allow_html=True)

    if not tasks:
        st.info("등록된 업무가 없습니다. 왼쪽 사이드바에서 새 업무를 추가해보세요!")
    else:
        # 카테고리 필터 박스
        categories = ["전체"] + list(set(t["category"] for t in tasks))
        selected_cat = st.selectbox("📂 구분별 필터", categories)
        
        filtered_tasks = tasks if selected_cat == "전체" else [t for t in tasks if t["category"] == selected_cat]

        for task in filtered_tasks:
            subtasks = task["subtasks"]
            task_id = task["id"]
            
            total_sub = len(subtasks)
            done_sub = sum(1 for s in subtasks if isinstance(s, dict) and s.get("done"))
            
            # 진행률 계산
            progress = (done_sub / total_sub) if total_sub > 0 else (1.0 if task["status"] == "완료" else 0.0)
            progress_percent = int(progress * 100)

            # 상태 자동 변경 로직
            if total_sub > 0:
                if progress_percent == 100:
                    new_status = "완료"
                    if not task.get("end_date"):
                        update_task_db(task_id, subtasks, new_status, end_date=str(date.today()))
                        task["end_date"] = str(date.today())
                elif progress_percent > 0:
                    new_status = "진행중"
                else:
                    new_status = "대기"
            else:
                new_status = task["status"]

            if new_status != task["status"]:
                task["status"] = new_status
                update_task_db(task_id, subtasks, new_status)

            start_str = task.get("start_date") or "미지정"
            due_str = task.get("due_date") or "미지정"
            end_str = task.get("end_date") or "-"

            # 메인 화면 카드 UI
            with st.container():
                st.markdown(f"""
                <div style="padding: 15px; border-radius: 8px; border: 1px solid #e0e0e0; margin-bottom: 5px; background-color: #ffffff;">
                    <span style="background-color: #f0f2f6; padding: 3px 8px; border-radius: 4px; font-size: 12px; font-weight: bold;">{task['category']}</span>
                    <span style="font-size: 16px; font-weight: bold; margin-left: 10px;">{task['title']}</span>
                    <span style="color: #666; font-size: 14px; margin-left: 10px;">({task['status']} - {progress_percent}%)</span>
                    <div style="margin-top: 8px; font-size: 13px; color: #555;">
                        📅 <b>전체 시작일:</b> {start_str} &nbsp;|&nbsp; ⏳ <b>전체 마감예정일:</b> {due_str} &nbsp;|&nbsp; ✅ <b>전체 종료일:</b> {end_str}
                    </div>
                </div>
                """, unsafe_allow_html=True)
                
                col_bar, col_btn1, col_btn2 = st.columns([4, 1, 1])
                with col_bar:
                    st.progress(progress)
                with col_btn1:
                    if st.button("⚙️ 세부업무 & 일정 편집", key=f"edit_btn_{task_id}", use_container_width=True):
                        st.session_state.editing_task_id = task_id
                        st.rerun()
                with col_btn2:
                    if st.button("🗑️ 전체 삭제", key=f"del_btn_{task_id}", use_container_width=True):
                        with st.spinner("삭제 중..."):
                            delete_task_db(task_id)
                        st.success("삭제되었습니다!")
                        st.rerun()
                st.markdown("---")

with tab2:
    st.subheader("📌 전체 업무 통계 대시보드")
    if tasks:
        tot_count = len(tasks)
        completed_count = sum(1 for t in tasks if t["status"] == "완료")
        in_progress_count = sum(1 for t in tasks if t["status"] == "진행중")
        pending_count = sum(1 for t in tasks if t["status"] == "대기")

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("전체 업무", f"{tot_count}개")
        m2.metric("진행 중", f"{in_progress_count}개")
        m3.metric("완료됨", f"{completed_count}개")
        m4.metric("대기 중", f"{pending_count}개")

        st.markdown("---")
        st.subheader("📂 업무 구분별 요약")
        cat_map = {}
        for t in tasks:
            cat = t["category"]
            if cat not in cat_map:
                cat_map[cat] = {"total": 0, "completed": 0}
            cat_map[cat]["total"] += 1
            if t["status"] == "완료":
                cat_map[cat]["completed"] += 1

        for cat, val in cat_map.items():
            rate = int((val["completed"] / val["total"]) * 100) if val["total"] > 0 else 0
            st.write(f"**{cat}** ({val['completed']}/{val['total']} 완료)")
            st.progress(rate / 100)

        st.markdown("---")
        st.subheader("📥 전체 업무 리스트 엑셀 다운로드")
        st.caption("다운로드할 업무 상태를 체크박스로 선택하고, 세부 업무별 일정 및 셀 병합이 적용된 엑셀 파일을 다운로드하세요.")

        # 상태 선택 체크박스 레이아웃 (가로 배치)
        st.write("**다운로드할 업무 상태 선택**")
        chk_col1, chk_col2, chk_col3 = st.columns(3)
        with chk_col1:
            sel_pending = st.checkbox("대기", value=True)
        with chk_col2:
            sel_progress = st.checkbox("진행중", value=True)
        with chk_col3:
            sel_completed = st.checkbox("완료", value=True)

        selected_statuses = []
        if sel_pending: selected_statuses.append("대기")
        if sel_progress: selected_statuses.append("진행중")
        if sel_completed: selected_statuses.append("완료")

        filtered_export_tasks = [t for t in tasks if t.get("status") in selected_statuses]

        if not filtered_export_tasks:
            st.warning("선택된 상태의 업무가 없습니다. 적어도 하나의 상태를 체크해주세요.")
        else:
            excel_data = []
            for t in filtered_export_tasks:
                subtasks = t["subtasks"]
                total_sub = len(subtasks)
                done_sub = sum(1 for s in subtasks if isinstance(s, dict) and s.get("done"))
                progress_percent = int((done_sub / total_sub) * 100) if total_sub > 0 else (100 if t["status"] == "완료" else 0)
                
                # 세부 업무 항목별로 각각 별도의 행(Row) 생성 및 세부 일정 포함
                if subtasks:
                    for s in subtasks:
                        if isinstance(s, dict):
                            sub_name = s.get('name', '')
                            sub_done = "완료" if s.get('done') else "진행중"
                            sub_memo = s.get('memo', '')
                            sub_sdate = s.get('start_date', '')
                            sub_edate = s.get('end_date', '')
                        else:
                            sub_name = str(s)
                            sub_done = "진행중"
                            sub_memo = ""
                            sub_sdate = ""
                            sub_edate = ""

                        excel_data.append({
                            "ID": t["id"],
                            "업무 구분": t["category"],
                            "업무 제목": t["title"],
                            "전체 상태": t["status"],
                            "전체 진행률(%)": f"{progress_percent}%",
                            "전체 시작일": t.get("start_date", ""),
                            "전체 마감예정일": t.get("due_date", ""),
                            "전체 종료일": t.get("end_date", ""),
                            "세부 업무명": sub_name,
                            "세부 완료 여부": sub_done,
                            "세부 시작일": sub_sdate,
                            "세부 종료일": sub_edate,
                            "추진계획 메모": sub_memo
                        })
                else:
                    excel_data.append({
                        "ID": t["id"],
                        "업무 구분": t["category"],
                        "업무 제목": t["title"],
                        "전체 상태": t["status"],
                        "전체 진행률(%)": f"{progress_percent}%",
                        "전체 시작일": t.get("start_date", ""),
                        "전체 마감예정일": t.get("due_date", ""),
                        "전체 종료일": t.get("end_date", ""),
                        "세부 업무명": "-",
                        "세부 완료 여부": "-",
                        "세부 시작일": "",
                        "세부 종료일": "",
                        "추진계획 메모": ""
                    })

            df = pd.DataFrame(excel_data)

            # Excel 파일 생성 및 스타일링 (셀 병합, 줄바꿈, 정렬)
            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name="업무진행상황")

            output.seek(0)
            wb = openpyxl.load_workbook(output)
            ws = wb.active
            
            # 셀 스타일 설정 (모든 셀 줄바꿈 및 정렬)
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center", horizontal="center" if cell.column in [1, 4, 5, 6, 7, 8, 10, 11, 12] else "left")

            # ID 기준(1열)으로 동일한 상위 업무 행들을 찾아 셀 병합 실행 (열 1~8 병합)
            merge_cols = [1, 2, 3, 4, 5, 6, 7, 8]
            start_row = 2
            max_row = ws.max_row
            
            while start_row <= max_row:
                current_id = ws.cell(row=start_row, column=1).value
                end_row = start_row
                
                while end_row + 1 <= max_row and ws.cell(row=end_row + 1, column=1).value == current_id:
                    end_row += 1
                
                if end_row > start_row:
                    for col in merge_cols:
                        ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                
                start_row = end_row + 1

            # 테두리 설정
            thin_border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style='thin', color='CCCCCC'),
                right=openpyxl.styles.Side(style='thin', color='CCCCCC'),
                top=openpyxl.styles.Side(style='thin', color='CCCCCC'),
                bottom=openpyxl.styles.Side(style='thin', color='CCCCCC')
            )
            
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.border = thin_border

            final_output = BytesIO()
            wb.save(final_output)
            excel_file = final_output.getvalue()

            today_str = datetime.now().strftime("%Y%m%d")
            st.download_button(
                label="📊 세부 일정 및 셀 병합된 엑셀 파일 다운로드 (.xlsx)",
                data=excel_file,
                file_name=f"업무진행상황_세부일정포함_{today_str}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    else:
        st.info("등록된 데이터가 없습니다.")
